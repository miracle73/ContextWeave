"""Extraction of CVs, web pages and GitHub repos. All output is untrusted text."""

import asyncio
import io
import ipaddress
import re
import socket
from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse

import httpx

from .config import settings

MAX_TEXT = 150_000
MAX_FETCH = 2_000_000


class IngestError(Exception):
    pass


def clean(text: str) -> str:
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT]


def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return clean("\n\n".join((p.extract_text() or "") for p in reader.pages))
    except Exception as e:  # noqa: BLE001
        raise IngestError(f"Could not read PDF: {e}") from e


def extract_docx(data: bytes) -> str:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise IngestError(f"Could not read DOCX: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return clean("\n\n".join(parts))


def extract_cv(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        text = extract_pdf(data)
    elif name.endswith(".docx"):
        text = extract_docx(data)
    else:
        raise IngestError("Only PDF and DOCX CVs are supported")
    if not text:
        raise IngestError("No text found (scanned PDFs need OCR first)")
    return text


class _TextParser(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "template", "head"}
    BLOCK = {"p", "div", "li", "br", "h1", "h2", "h3", "h4", "section", "article", "tr"}

    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self.skip = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n\n")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> tuple[str, str]:
    p = _TextParser()
    p.feed(html)
    return p.title.strip(), clean("".join(p.out))


async def assert_public_url(url: str) -> None:
    """SSRF guard: only http(s) to public IPs."""
    u = urlparse(url)
    if u.scheme not in ("http", "https") or not u.hostname:
        raise IngestError("URL must be http(s)")
    try:
        infos = await asyncio.get_running_loop().getaddrinfo(u.hostname, None)
    except socket.gaierror as e:
        raise IngestError(f"Cannot resolve {u.hostname}") from e
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global:
            raise IngestError("URL resolves to a private or reserved address")


async def _get(client: httpx.AsyncClient, url: str, headers: dict | None = None) -> httpx.Response:
    for _ in range(4):
        await assert_public_url(url)
        async with client.stream("GET", url, headers=headers) as r:
            if r.is_redirect and "location" in r.headers:
                url = urljoin(url, r.headers["location"])
                continue
            body = b""
            async for chunk in r.aiter_bytes():
                body += chunk
                if len(body) > MAX_FETCH:
                    break
            return httpx.Response(r.status_code, headers=r.headers, content=body, request=r.request)
    raise IngestError("Too many redirects")


async def fetch_url(url: str) -> tuple[str, str]:
    async with httpx.AsyncClient(timeout=15, follow_redirects=False,
                                 headers={"User-Agent": "ContextWeave/1.0"}) as c:
        r = await _get(c, url)
    if r.status_code >= 400:
        raise IngestError(f"Fetch failed with HTTP {r.status_code}")
    ctype = r.headers.get("content-type", "")
    raw = r.content.decode(r.encoding or "utf-8", errors="replace")
    if "html" in ctype:
        title, text = html_to_text(raw)
    elif ctype.startswith("text/") or "json" in ctype:
        title, text = "", clean(raw)
    else:
        raise IngestError(f"Unsupported content type: {ctype or 'unknown'}")
    if not text:
        raise IngestError("No readable text (JS-rendered pages are not supported)")
    return title or url, text


_GH = re.compile(r"^(?:https?://github\.com/)?([\w.-]+)/([\w.-]+?)(?:\.git)?/?$")
KEY_FILES = ("package.json", "pyproject.toml", "requirements.txt", "go.mod", "Cargo.toml",
             "pom.xml", "Dockerfile", "docker-compose.yml")


def parse_repo(ref: str) -> tuple[str, str]:
    m = _GH.match(ref.strip())
    if not m:
        raise IngestError("Use owner/repo or https://github.com/owner/repo")
    return m.group(1), m.group(2)


async def fetch_github(ref: str) -> tuple[str, str]:
    owner, repo = parse_repo(ref)
    headers = {"Accept": "application/vnd.github+json", "X-GitHub-Api-Version": "2022-11-28"}
    if settings.github_token:
        headers["Authorization"] = f"Bearer {settings.github_token}"
    base = f"https://api.github.com/repos/{owner}/{repo}"
    async with httpx.AsyncClient(timeout=20, headers=headers) as c:
        r = await c.get(base)
        if r.status_code == 404:
            raise IngestError("Repository not found or private")
        if r.status_code in (403, 429):
            raise IngestError("GitHub rate limit hit; set GITHUB_TOKEN")
        r.raise_for_status()
        meta = r.json()
        branch = meta.get("default_branch", "main")
        raw_hdr = {"Accept": "application/vnd.github.raw+json"}
        readme, langs, tree = await asyncio.gather(
            c.get(f"{base}/readme", headers=raw_hdr), c.get(f"{base}/languages"),
            c.get(f"{base}/git/trees/{branch}?recursive=1"))
        paths = [t["path"] for t in tree.json().get("tree", []) if t.get("type") == "blob"] if tree.is_success else []
        parts = [
            f"GitHub repository {owner}/{repo}",
            f"Description: {meta.get('description') or 'n/a'}",
            f"Topics: {', '.join(meta.get('topics', [])) or 'n/a'}",
            f"Languages: {', '.join(langs.json().keys()) if langs.is_success else 'n/a'}",
            f"Stars: {meta.get('stargazers_count', 0)}",
            "File tree (partial):\n" + "\n".join(paths[:200]),
        ]
        if readme.is_success:
            parts.append("README:\n" + readme.text[:40_000])
        for f in [p for p in paths if p.split("/")[-1] in KEY_FILES and p.count("/") <= 1][:6]:
            fr = await c.get(f"{base}/contents/{f}?ref={branch}", headers=raw_hdr)
            if fr.is_success:
                parts.append(f"File {f}:\n{fr.text[:6000]}")
    return f"{owner}/{repo}", clean("\n\n".join(parts))
